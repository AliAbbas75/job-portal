"""Resume upload + parse (T-050) and resume PDF export (T-052), master flow §4.5.

The uploaded resume is kept as the candidate's "resume" document (supporting material). Parsing
only suggests values; the candidate confirms each section, which saves it to the profile.
"""

from datetime import date
from io import BytesIO

from docx import Document as DocxDocument
from flask import current_app
from fpdf import FPDF
from pypdf import PdfReader

from app.services import document_service
from app.services.resume_parser_service import parse_resume_text
from app.services.storage_service import DOCX, RESUME_TYPES, path_of
from app.utils.dates import age_on


def extract_text(path, content_type):
    """Plain text of a PDF or DOCX file; '' when none can be read (e.g. a scanned image)."""
    try:
        if content_type == DOCX:
            document = DocxDocument(str(path))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(lines)
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages[:10])
    except Exception as error:  # a damaged file must not break the upload
        current_app.logger.warning("Could not read resume text: %s", type(error).__name__)
        return ""


def upload_and_parse(account, profile, file):
    """Stores the resume and returns {documentId, parsed, lowConfidence, sections}."""
    document = document_service.upload(account, profile, "resume", file, RESUME_TYPES)
    text = extract_text(path_of(document.storage_key), document.content_type)
    return {"documentId": document.id, **parse_resume_text(text)}


# ---- PDF export ----


def _latin(text):
    # The built-in PDF fonts cover Latin-1 only; other characters print as "?".
    return str(text or "").encode("latin-1", "replace").decode("latin-1")


class _ResumePdf(FPDF):
    def heading(self, text):
        self.ln(3)
        self.set_font("helvetica", "B", 12)
        self.set_text_color(31, 77, 54)  # heritage green
        self.cell(0, 7, _latin(text), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(31, 77, 54)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(2)
        self.set_text_color(0, 0, 0)

    def line_text(self, text, bold=False, size=10):
        self.set_font("helvetica", "B" if bold else "", size)
        self.multi_cell(0, 5, _latin(text), new_x="LMARGIN", new_y="NEXT")


def build_resume_pdf(profile, qualification_names):
    """A one-column resume of the profile. Returns PDF bytes."""
    pdf = _ResumePdf(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.add_page()

    pdf.set_font("helvetica", "B", 18)
    pdf.cell(0, 9, _latin(profile.full_name or "Candidate"), new_x="LMARGIN", new_y="NEXT")
    contact = [profile.account.mobile, profile.email, profile.current_address]
    pdf.line_text(" | ".join(_latin(c) for c in contact if c), size=9)

    personal = []
    if profile.father_name:
        personal.append(f"Father's name: {profile.father_name}")
    if profile.dob:
        personal.append(f"Date of birth: {profile.dob.strftime('%d %b %Y')} "
                        f"(age {age_on(profile.dob, date.today())})")  # fmt: skip
    if profile.domicile_district:
        personal.append(
            f"Domicile: {profile.domicile_district.name}, {profile.domicile_province_code}"
        )
    if personal:
        pdf.heading("Personal details")
        for item in personal:
            pdf.line_text(item)

    if profile.statement_of_purpose:
        pdf.heading("Statement of purpose")
        pdf.line_text(profile.statement_of_purpose)

    if profile.education:
        pdf.heading("Education")
        for e in sorted(profile.education, key=lambda e: e.passing_year, reverse=True):
            level = qualification_names.get(e.qualification_level_code, e.qualification_level_code)
            pdf.line_text(f"{level} in {e.discipline} ({e.passing_year})", bold=True)
            pdf.line_text(f"{e.institution} - {float(e.marks_percent):g}%")

    if profile.experience:
        pdf.heading("Experience")
        for x in sorted(profile.experience, key=lambda x: x.start_date, reverse=True):
            end = "present" if x.is_current else x.end_date.strftime("%b %Y")
            pdf.line_text(f"{x.designation}, {x.organization}", bold=True)
            pdf.line_text(f"{x.start_date.strftime('%b %Y')} - {end}")

    if profile.registrations:
        pdf.heading("Professional registration")
        for r in profile.registrations:
            until = f", valid until {r.valid_until.strftime('%d %b %Y')}" if r.valid_until else ""
            pdf.line_text(f"{r.body}: {r.registration_no}{until}")

    if profile.publications:
        pdf.heading("Publications")
        for p in profile.publications:
            details = ", ".join(str(v) for v in (p.venue, p.year) if v)
            pdf.line_text(f"{p.title}" + (f" ({details})" if details else ""))

    if profile.skills:
        pdf.heading("Skills")
        pdf.line_text(", ".join(profile.skills))

    if profile.references:
        pdf.heading("References")
        for ref in profile.references:
            parts = [ref.name, ref.designation, ref.organization, ref.phone, ref.email]
            pdf.line_text(", ".join(p for p in parts if p))

    buffer = BytesIO()
    pdf.output(buffer)
    return buffer.getvalue()
