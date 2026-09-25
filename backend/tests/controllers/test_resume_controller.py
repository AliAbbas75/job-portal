import io
from datetime import date

from docx import Document as DocxDocument
from fpdf import FPDF

from app.extensions import db
from app.models import Document, EducationRecord

LINES = [
    "Sara Ahmed",
    "Email: sara@example.com",
    "Education",
    "BS Computer Science from National University 2020 81%",
    "Experience",
    "Feb 2021 - Present Software Engineer at Railway IT Directorate",
]


def _docx():
    document = DocxDocument()
    for line in LINES:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    for line in LINES:
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _upload(client, headers, data, name):
    return client.post(
        "/api/resume",
        data={"file": (io.BytesIO(data), name)},
        headers=headers,
        content_type="multipart/form-data",
    )


def test_upload_docx_and_pdf_resumes(client, make_candidate, auth_headers):
    headers = auth_headers(make_candidate())
    for data, name in [(_docx(), "cv.docx"), (_pdf(), "cv.pdf")]:
        response = _upload(client, headers, data, name)
        assert response.status_code == 201, response.get_json()
        body = response.get_json()
        assert body["parsed"] is True
        assert body["sections"]["contact"]["email"]["value"] == "sara@example.com"
        assert body["sections"]["education"][0]["level"] == "bachelor16"
        assert body["sections"]["experience"][0]["current"] is True
    # Stored as the (single current) resume document; nothing written to the profile.
    current = db.session.query(Document).filter_by(document_type_code="resume", archived_at=None)
    assert current.count() == 1
    assert db.session.query(EducationRecord).count() == 0


def test_resume_accepts_only_pdf_or_word(client, make_candidate, auth_headers):
    headers = auth_headers(make_candidate())
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    response = _upload(client, headers, png, "cv.png")
    assert response.status_code == 422 and response.get_json()["code"] == "file_type_not_allowed"
    # ...and Word files are still refused for ordinary documents.
    response = client.post(
        "/api/documents",
        data={"type": "degree", "file": (io.BytesIO(_docx()), "degree.docx")},
        headers=headers,
        content_type="multipart/form-data",
    )
    assert response.status_code == 422


def test_scanned_resume_is_not_parsed(client, make_candidate, auth_headers):
    blank = FPDF()
    blank.add_page()  # no text, like a scanned image
    response = _upload(client, auth_headers(make_candidate()), bytes(blank.output()), "scan.pdf")
    assert response.status_code == 201 and response.get_json()["parsed"] is False


def test_bps15_sections_and_resume_pdf(client, make_candidate, auth_headers):
    account = make_candidate()
    headers = auth_headers(account)
    profile = account.profile
    profile.full_name = "Sara Ahmed"
    profile.dob = date(1995, 5, 1)
    db.session.commit()

    assert (
        client.put(
            "/api/profile/statement",
            json={"statementOfPurpose": "I build rail systems."},
            headers=headers,
        ).get_json()["statementOfPurpose"]
        == "I build rail systems."
    )
    body = client.post(
        "/api/profile/registrations",
        json={"body": "PEC", "registrationNo": "CIVIL/1", "validUntil": "2030-01-01"},
        headers=headers,
    ).get_json()
    assert body["registrations"][0]["registrationNo"] == "CIVIL/1"
    body = client.post(
        "/api/profile/publications", json={"title": "Rail wear", "year": 2021}, headers=headers
    ).get_json()
    assert body["publications"][0]["title"] == "Rail wear"
    body = client.post(
        "/api/profile/references",
        json={"name": "Dr. Test", "email": "ref@example.com"},
        headers=headers,
    ).get_json()
    reference_id = body["references"][0]["id"]
    body = client.delete(f"/api/profile/references/{reference_id}", headers=headers).get_json()
    assert body["references"] == []

    pdf = client.get("/api/resume/pdf", headers=headers)
    assert pdf.status_code == 200 and pdf.mimetype == "application/pdf"
    assert pdf.data.startswith(b"%PDF")


def test_jobs_say_which_tier_they_use(client, publish_job):
    publish_job(bps=11)
    publish_job(bps=17, title="Officer")
    tiers = {j["bps"]: j["resumeTier"] for j in client.get("/api/jobs").get_json()["items"]}
    assert tiers == {11: False, 17: True}
